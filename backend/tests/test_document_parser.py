import json
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook  # type: ignore[import-untyped]

from app.workers.document_parser import ParserError, parse_document
from app.workers.document_security import run_sandboxed_parser


def test_docx_parser_extracts_text_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "quote.docx"
    document = WordDocument()
    document.add_paragraph("Supplier quotation")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Laptop"
    table.cell(1, 1).text = "1000"
    document.save(path)

    result = parse_document(
        path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    assert result["kind"] == "native"
    assert result["pages"][0]["text"] == "Supplier quotation"  # type: ignore[index]
    assert result["pages"][0]["tables"][0]["rows"][1] == [  # type: ignore[index]
        "Laptop",
        "1000",
    ]


def test_xlsx_parser_preserves_sheet_rows(tmp_path: Path) -> None:
    path = tmp_path / "quote.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pricing"
    sheet.append(["Item", "Price"])
    sheet.append(["Laptop", 1000])
    workbook.save(path)

    result = parse_document(
        path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    page = result["pages"][0]  # type: ignore[index]
    assert page["source_label"] == "Pricing"
    assert page["text"] == "Item\tPrice\nLaptop\t1000"


def test_office_parser_rejects_invalid_package(tmp_path: Path) -> None:
    path = tmp_path / "fake.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("unexpected.txt", "not an office package")

    with pytest.raises(ParserError, match="package structure"):
        parse_document(
            path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )


def test_declared_media_type_must_match_file_signature(tmp_path: Path) -> None:
    path = tmp_path / "fake.pdf"
    path.write_bytes(b"not a pdf")

    with pytest.raises(ParserError, match="signature"):
        parse_document(path, "application/pdf")


def test_parser_subprocess_returns_bounded_json(tmp_path: Path) -> None:
    path = tmp_path / "input.bin"
    path.write_bytes(b"data")
    expected = {"parser": "test", "parser_version": "1", "kind": "native", "pages": []}

    result = run_sandboxed_parser(
        path,
        "application/pdf",
        parser_command=[sys.executable, "-c", f"import json; print(json.dumps({expected!r}))"],
        timeout_seconds=5,
        memory_bytes=256 * 1024 * 1024,
        output_bytes=4096,
    )

    assert json.dumps(result, sort_keys=True) == json.dumps(expected, sort_keys=True)


def test_parser_subprocess_cannot_create_network_sockets(tmp_path: Path) -> None:
    path = tmp_path / "input.bin"
    path.write_bytes(b"data")
    script = (
        "import json,socket; "
        "\ntry: socket.socket()"
        "\nexcept PermissionError: print(json.dumps({'network_blocked': True}))"
        "\nelse: print(json.dumps({'network_blocked': False}))"
    )

    result = run_sandboxed_parser(
        path,
        "application/pdf",
        parser_command=[sys.executable, "-c", script],
        timeout_seconds=5,
        memory_bytes=256 * 1024 * 1024,
        output_bytes=4096,
    )

    assert result == {"network_blocked": True}
