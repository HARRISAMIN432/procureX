import argparse
import json
import subprocess
import tempfile
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from openpyxl import load_workbook  # type: ignore[import-untyped]
from PIL import Image
from pypdf import PdfReader

PARSER_NAME = "procurex-safe-parser"
PARSER_VERSION = "1"
MAX_PAGES = 500
MAX_PAGE_CHARACTERS = 2_000_000
MAX_ARCHIVE_ENTRIES = 10_000
MAX_ARCHIVE_EXPANDED_BYTES = 250 * 1024 * 1024
MAX_SHEET_ROWS = 100_000
MAX_SHEET_COLUMNS = 500


class ParserError(RuntimeError):
    pass


@dataclass(slots=True)
class ParsedPage:
    page_number: int
    text: str
    source_label: str | None = None
    width: float | None = None
    height: float | None = None
    ocr_confidence: float | None = None
    tables: list[dict[str, Any]] | None = None

    def as_payload(self) -> dict[str, Any]:
        value = asdict(self)
        value["tables"] = self.tables or []
        return value


def _require_magic(path: Path, media_type: str) -> None:
    prefix = path.read_bytes()[:8]
    signatures = {
        "application/pdf": prefix.startswith(b"%PDF-"),
        "image/png": prefix == b"\x89PNG\r\n\x1a\n",
        "image/jpeg": prefix.startswith(b"\xff\xd8\xff"),
    }
    if media_type in signatures and not signatures[media_type]:
        raise ParserError("File signature does not match the declared media type")
    if media_type.endswith(("wordprocessingml.document", "spreadsheetml.sheet")):
        if not prefix.startswith(b"PK\x03\x04"):
            raise ParserError("Office document is not a ZIP container")


def _validate_office_archive(path: Path, required_entry: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ParserError("Office archive contains too many entries")
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or required_entry not in names:
                raise ParserError("Office archive has an invalid package structure")
            expanded = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise ParserError("Encrypted Office archives are not supported")
                expanded += entry.file_size
                if expanded > MAX_ARCHIVE_EXPANDED_BYTES:
                    raise ParserError("Office archive expands beyond the safety limit")
                if entry.compress_size == 0 and entry.file_size:
                    raise ParserError("Office archive contains an invalid compressed entry")
                if entry.compress_size and entry.file_size / entry.compress_size > 1_000:
                    raise ParserError("Office archive compression ratio exceeds the safety limit")
    except (OSError, zipfile.BadZipFile) as exc:
        raise ParserError("Office archive is corrupt") from exc


def _bounded_text(value: str) -> str:
    if len(value) > MAX_PAGE_CHARACTERS:
        raise ParserError("Extracted page text exceeds the safety limit")
    return value


def _ocr_image(path: Path) -> str:
    try:
        result = subprocess.run(  # noqa: S603
            ["tesseract", str(path), "stdout"],
            capture_output=True,
            check=False,
            timeout=45,
            text=True,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise ParserError("OCR engine is unavailable or timed out") from exc
    if result.returncode != 0:
        raise ParserError("OCR engine failed")
    return _bounded_text(result.stdout)


def _parse_pdf(path: Path) -> tuple[list[ParsedPage], str]:
    try:
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ParserError("Encrypted PDFs are not supported")
        if not 1 <= len(reader.pages) <= MAX_PAGES:
            raise ParserError("PDF page count exceeds the safety limit")
        pages: list[ParsedPage] = []
        used_native = False
        used_ocr = False
        for index, page in enumerate(reader.pages, start=1):
            text = _bounded_text(page.extract_text() or "")
            if text.strip():
                used_native = True
            else:
                with tempfile.TemporaryDirectory(prefix="procurex-pdf-page-") as directory:
                    output = Path(directory) / "page"
                    result = subprocess.run(  # noqa: S603
                        [
                            "pdftoppm",
                            "-f",
                            str(index),
                            "-l",
                            str(index),
                            "-singlefile",
                            "-png",
                            "-r",
                            "200",
                            str(path),
                            str(output),
                        ],
                        capture_output=True,
                        check=False,
                        timeout=45,
                    )
                    if result.returncode != 0:
                        raise ParserError("PDF rendering failed")
                    text = _ocr_image(output.with_suffix(".png"))
                    used_ocr = True
            pages.append(
                ParsedPage(
                    page_number=index,
                    source_label=f"Page {index}",
                    text=text,
                    width=float(page.mediabox.width),
                    height=float(page.mediabox.height),
                )
            )
    except ParserError:
        raise
    except Exception as exc:
        raise ParserError("PDF parsing failed") from exc
    kind = "hybrid" if used_native and used_ocr else "ocr" if used_ocr else "native"
    return pages, kind


def _parse_docx(path: Path) -> list[ParsedPage]:
    _validate_office_archive(path, "word/document.xml")
    try:
        document = WordDocument(str(path))
        lines = [paragraph.text for paragraph in document.paragraphs]
        tables: list[dict[str, Any]] = []
        for table in document.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            tables.append({"rows": rows})
        text = _bounded_text("\n".join(lines))
        return [ParsedPage(page_number=1, source_label="Document", text=text, tables=tables)]
    except ParserError:
        raise
    except Exception as exc:
        raise ParserError("DOCX parsing failed") from exc


def _parse_xlsx(path: Path) -> list[ParsedPage]:
    _validate_office_archive(path, "xl/workbook.xml")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if not 1 <= len(workbook.worksheets) <= MAX_PAGES:
            raise ParserError("Workbook sheet count exceeds the safety limit")
        pages: list[ParsedPage] = []
        for page_number, sheet in enumerate(workbook.worksheets, start=1):
            if sheet.max_row > MAX_SHEET_ROWS or sheet.max_column > MAX_SHEET_COLUMNS:
                raise ParserError("Worksheet dimensions exceed the safety limit")
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(["" if value is None else str(value) for value in row])
            text = _bounded_text("\n".join("\t".join(row) for row in rows))
            pages.append(
                ParsedPage(
                    page_number=page_number,
                    source_label=sheet.title[:200],
                    text=text,
                    tables=[{"rows": rows}],
                )
            )
        workbook.close()
        return pages
    except ParserError:
        raise
    except Exception as exc:
        raise ParserError("XLSX parsing failed") from exc


def _parse_image(path: Path) -> list[ParsedPage]:
    try:
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            width, height = image.size
            if width * height > 100_000_000:
                raise ParserError("Image dimensions exceed the safety limit")
        return [
            ParsedPage(
                page_number=1,
                source_label="Image 1",
                text=_ocr_image(path),
                width=float(width),
                height=float(height),
            )
        ]
    except ParserError:
        raise
    except Exception as exc:
        raise ParserError("Image parsing failed") from exc


def parse_document(path: Path, media_type: str) -> dict[str, Any]:
    _require_magic(path, media_type)
    if media_type == "application/pdf":
        pages, kind = _parse_pdf(path)
    elif media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        pages, kind = _parse_docx(path), "native"
    elif media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        pages, kind = _parse_xlsx(path), "native"
    elif media_type in {"image/jpeg", "image/png"}:
        pages, kind = _parse_image(path), "ocr"
    else:
        raise ParserError("Unsupported document media type")
    if not pages:
        raise ParserError("Parser produced no pages")
    return {
        "parser": PARSER_NAME,
        "parser_version": PARSER_VERSION,
        "kind": kind,
        "pages": [page.as_payload() for page in pages],
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("media_type")
    arguments = parser.parse_args()
    try:
        print(json.dumps(parse_document(arguments.path, arguments.media_type)))
    except ParserError as exc:
        print(json.dumps({"error_code": "unsafe_or_unparseable", "error_detail": str(exc)}))
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
